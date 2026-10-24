import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from research_tool import load_records, summarize_records, write_excel


class ResearchToolTests(unittest.TestCase):
    def test_load_csv_and_build_summary(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            csv_path = Path(temp_dir) / "survey.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=["age", "opinion"])
                writer.writeheader()
                writer.writerow({"age": "21", "opinion": "good"})
                writer.writerow({"age": "19", "opinion": "good"})
                writer.writerow({"age": "20", "opinion": "great"})

            records = load_records(csv_path)
            summary = summarize_records(records)

        self.assertEqual(len(records), 3)
        self.assertEqual({row["field"] for row in summary}, {"age", "opinion"})
        age_summary = next(row for row in summary if row["field"] == "age")
        self.assertEqual(age_summary["responses"], 3)
        self.assertEqual(age_summary["average"], 20.0)
        self.assertEqual(age_summary["minimum"], 19.0)
        self.assertEqual(age_summary["maximum"], 21.0)

    def test_write_excel_with_pagination(self):
        records = [
            {"id": "1", "answer": "yes"},
            {"id": "2", "answer": "no"},
            {"id": "3", "answer": "yes"},
        ]
        summary = summarize_records(records)

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "report.xlsx"
            write_excel(records, summary, output_path, page_size=2)

            workbook = load_workbook(output_path)
            self.assertIn("RawData", workbook.sheetnames)
            self.assertIn("Summary", workbook.sheetnames)
            self.assertIn("Page_1", workbook.sheetnames)
            self.assertIn("Page_2", workbook.sheetnames)
            self.assertEqual(workbook["RawData"].max_row, 4)
            self.assertEqual(workbook["Page_1"].max_row, 3)
            self.assertEqual(workbook["Page_2"].max_row, 2)


if __name__ == "__main__":
    unittest.main()
