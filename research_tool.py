import argparse
import csv
import json
from pathlib import Path
from statistics import mean

from openpyxl import Workbook


def _coerce_number(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def load_records(input_path):
    path = Path(input_path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() == ".json":
        with path.open(encoding="utf-8") as handle:
            payload = json.load(handle)
            if not isinstance(payload, list):
                raise ValueError("JSON input must be a list of objects.")
            if payload and not isinstance(payload[0], dict):
                raise ValueError("JSON list entries must be objects.")
            return payload
    raise ValueError("Only CSV and JSON inputs are supported.")


def summarize_records(records):
    if not records:
        return []
    columns = sorted({key for row in records for key in row.keys()})
    summary = []
    for column in columns:
        values = [row.get(column) for row in records]
        non_empty = [value for value in values if str(value).strip() != "" and value is not None]
        numeric_values = [_coerce_number(value) for value in non_empty]
        numeric_values = [value for value in numeric_values if value is not None]

        row = {
            "field": column,
            "responses": len(non_empty),
            "unique_values": len(set(map(str, non_empty))),
            "average": "",
            "minimum": "",
            "maximum": "",
        }
        if numeric_values:
            row["average"] = round(mean(numeric_values), 4)
            row["minimum"] = min(numeric_values)
            row["maximum"] = max(numeric_values)
        summary.append(row)
    return summary


def write_excel(records, summary, output_path, page_size=500):
    workbook = Workbook()
    raw_sheet = workbook.active
    raw_sheet.title = "RawData"

    columns = sorted({key for row in records for key in row.keys()}) if records else []
    if columns:
        raw_sheet.append(columns)
        for row in records:
            raw_sheet.append([row.get(column, "") for column in columns])

    summary_sheet = workbook.create_sheet("Summary")
    summary_columns = ["field", "responses", "unique_values", "average", "minimum", "maximum"]
    summary_sheet.append(summary_columns)
    for row in summary:
        summary_sheet.append([row[column] for column in summary_columns])

    if records and page_size > 0:
        for index in range(0, len(records), page_size):
            page_sheet = workbook.create_sheet(f"Page_{(index // page_size) + 1}")
            page_sheet.append(columns)
            for row in records[index : index + page_size]:
                page_sheet.append([row.get(column, "") for column in columns])

    workbook.save(output_path)


def build_parser():
    parser = argparse.ArgumentParser(
        description="Convert research/survey datasets to paginated Excel output with summary metrics."
    )
    parser.add_argument("input", help="Input data file (.csv or .json)")
    parser.add_argument("output", help="Output Excel file (.xlsx)")
    parser.add_argument(
        "--page-size",
        type=int,
        default=500,
        help="Rows per pagination sheet (default: 500)",
    )
    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()
    records = load_records(args.input)
    summary = summarize_records(records)
    write_excel(records, summary, args.output, page_size=args.page_size)
    print(f"Generated {args.output} with {len(records)} records.")


if __name__ == "__main__":
    main()
